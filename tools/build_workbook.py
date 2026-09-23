# Generates the formatted SDG Hackathon evaluation workbook.
# Usage: python build_workbook.py OUTPUT.xlsx [demo]   (needs: pip install openpyxl)
import sys, random
from openpyxl import Workbook
from openpyxl.styles import Protection, Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

DEMO = len(sys.argv) > 2 and sys.argv[2] == "demo"
OUT = sys.argv[1]
random.seed(7)

F = "Arial"
def font(**k): return Font(name=F, **k)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
NAVY, WHITE = "1F3864", "FFFFFF"
INPUT = "FFF2CC"; BLUE = "0000FF"
E_FILLS = ["DDEBF7", "E2EFDA", "FCE4D6"]
AVGF, RESF, ATTF, SDGF = "EDEDED", "D9E1F2", "E4DFEC", "E2F0D9"

def hdr(c, text, bg=NAVY, fg=WHITE, bold=True, size=10):
    c.value = text; c.font = font(bold=bold, color=fg, size=size)
    c.fill = fill(bg); c.alignment = CENTER; c.border = BOX

def box(ws, rng):
    for row in ws[rng]:
        for c in row: c.border = BOX

def inp(c, v=None, align=CENTER):
    if v is not None: c.value = v
    c.fill = fill(INPUT); c.font = font(color=BLUE, size=10)
    c.alignment = align; c.border = BOX; c.protection = Protection(locked=False)

def calc(c, v, fmt=None, bg=None, bold=False, align=CENTER):
    c.value = v; c.font = font(size=10, bold=bold); c.alignment = align; c.border = BOX
    if fmt: c.number_format = fmt
    if bg: c.fill = fill(bg)

def unhide_axes(ch):
    ch.x_axis.delete = False; ch.y_axis.delete = False

# ---------------------------------------------------------------- reference data
POS = [("PO1", "Engineering Knowledge"), ("PO2", "Problem Analysis"),
       ("PO3", "Design/Development of Solutions"), ("PO4", "Conduct Investigations of Complex Problems"),
       ("PO5", "Engineering Tool Usage"), ("PO6", "The Engineer and The World"), ("PO7", "Ethics"),
       ("PO8", "Individual and Collaborative Team Work"), ("PO9", "Communication"),
       ("PO10", "Project Management and Finance"), ("PO11", "Life-Long Learning"),
       ("PSO1", "Intelligent Systems: Select appropriate technologies to analyse, design, implement, and deployment of smart and intelligent systems"),
       ("PSO2", "Contemporary Systems: Design, and development of efficient IT solutions for challenging issues through experiential learning")]
OUTCOMES = [p[0] for p in POS]   # 13 PO/PSO attainment columns (SDG is rated separately)
NO = len(OUTCOMES)
SDGS = ["SDG 1 - No Poverty", "SDG 2 - Zero Hunger", "SDG 3 - Good Health and Well-being",
        "SDG 4 - Quality Education", "SDG 5 - Gender Equality", "SDG 6 - Clean Water and Sanitation",
        "SDG 7 - Affordable and Clean Energy", "SDG 8 - Decent Work and Economic Growth",
        "SDG 9 - Industry, Innovation and Infrastructure", "SDG 10 - Reduced Inequalities",
        "SDG 11 - Sustainable Cities and Communities", "SDG 12 - Responsible Consumption and Production",
        "SDG 13 - Climate Action", "SDG 14 - Life Below Water", "SDG 15 - Life on Land",
        "SDG 16 - Peace, Justice and Strong Institutions", "SDG 17 - Partnerships for the Goals"]

# criterion, short, max, excellent, good, satisfactory, needs-improvement
CRIT = [
 ("Problem Identification & SDG Relevance", "C1 Problem & SDG", 10,
  "Real, well-scoped problem explicitly linked to specific SDG targets (e.g., 6.1, 13.1); stakeholders and measurable impact indicators identified with evidence.",
  "Relevant problem mapped to an SDG goal; scope and stakeholders mostly clear; limited supporting evidence.",
  "Problem stated in general terms; SDG link is broad or superficial; scope vague.",
  "Problem unclear or trivial; no credible SDG connection."),
 ("Research & Analysis of Existing Solutions", "C2 Research & Analysis", 10,
  "Thorough review of existing solutions, data and literature; gaps clearly articulated and used to justify the proposal.",
  "Reviews key existing solutions and identifies some gaps.",
  "Limited survey; gaps asserted without evidence.",
  "No analysis of existing work or data."),
 ("Solution Design & Innovation", "C3 Design & Innovation", 15,
  "Original, well-architected solution (architecture/flow diagrams, modules, data flow); clearly innovative beyond existing approaches.",
  "Sound design with some novelty; architecture mostly complete.",
  "Conventional design with gaps; minimal innovation.",
  "No coherent design, or a copy of an existing solution."),
 ("Technical Implementation & Tool Usage", "C4 Implementation", 20,
  "Working prototype demonstrates core features end-to-end; modern tools/frameworks used effectively; clean, version-controlled code.",
  "Prototype works for most core features; appropriate tools used.",
  "Partial prototype or mock-up with limited functionality.",
  "No working implementation."),
 ("Intelligent & Smart System Features (AI/ML, IoT, Analytics)", "C5 Intelligent Features", 10,
  "Meaningful use of AI/ML, data analytics, IoT or automation with justified technology/model choice and measured performance.",
  "Intelligent component present and working; limited evaluation.",
  "Intelligent feature only proposed or rudimentary (hard-coded rules).",
  "No intelligent or data-driven element."),
 ("Sustainability, Societal Impact & Ethics", "C6 Impact & Ethics", 10,
  "Quantified social/environmental impact; addresses privacy, safety, inclusivity, bias and environmental footprint.",
  "Impact and ethical aspects discussed with some specifics.",
  "Impact claimed in general terms; ethics mentioned only briefly.",
  "Impact and ethical considerations ignored."),
 ("Feasibility, Scalability & Cost", "C7 Feasibility & Cost", 10,
  "Realistic deployment plan, cost estimate or business model, scalability path and risks identified.",
  "Feasible, with rough costing and some scalability considerations.",
  "Feasibility asserted; costing or scaling not addressed.",
  "Infeasible, or no deployment thinking."),
 ("Teamwork & Project Management", "C8 Teamwork", 5,
  "Clear roles and balanced contribution; planned timeline/task tracking evident; every member answers confidently.",
  "Roles defined; most members contribute.",
  "Uneven contribution; little evidence of planning.",
  "One or two members did most of the work; no planning."),
 ("Presentation, Demo & Q&A", "C9 Presentation & Q&A", 10,
  "Clear, time-bound pitch; compelling live demo; accurate, confident answers to technical questions.",
  "Clear presentation and demo; answers most questions.",
  "Presentation unclear or over time; weak answers.",
  "Disorganised; unable to answer questions."),
]
NC = len(CRIT)
#            PO1..PO11                    PSO1 PSO2 SDG
MAP = [
 [1,3,0,1,0,3,0,0,0,0,1, 0,1, 3],
 [2,3,0,3,1,1,0,0,0,0,2, 0,0, 1],
 [2,2,3,1,0,1,0,0,0,0,0, 2,3, 1],
 [2,0,2,0,3,0,0,1,0,0,1, 3,2, 0],
 [2,1,1,2,2,0,0,0,0,0,1, 3,1, 0],
 [0,0,1,0,0,3,3,0,0,0,0, 0,1, 3],
 [0,0,1,0,0,2,0,0,0,3,0, 0,2, 2],
 [0,0,0,0,0,0,1,3,1,2,0, 0,0, 0],
 [0,0,0,0,0,0,0,2,3,0,1, 0,1, 0],
]

MAP = [row[:NO] for row in MAP]

wb = Workbook()
wb.remove(wb.active)
ws_ins = wb.create_sheet("Instructions")
ws_dash = wb.create_sheet("Dashboard")
ws_rub = wb.create_sheet("Rubric")
ws_map = wb.create_sheet("Mapping")
EVENTS = [  # sheet, title, kind, source sheets
 ("3A Intra", "3rd Semester Section A — Intra-Class Hackathon", "intra", "3rd Sem - A"),
 ("3B Intra", "3rd Semester Section B — Intra-Class Hackathon", "intra", "3rd Sem - B"),
 ("3 Inter", "3rd Semester — Inter-Section Hackathon (Finals)", "inter", ("3A Intra", "3B Intra", "3rd Sem - A", "3rd Sem - B")),
 ("5A Intra", "5th Semester Section A — Intra-Class Hackathon", "intra", "5th Sem - A"),
 ("5B Intra", "5th Semester Section B — Intra-Class Hackathon", "intra", "5th Sem - B"),
 ("5 Inter", "5th Semester — Inter-Section Hackathon (Finals)", "inter", ("5A Intra", "5B Intra", "5th Sem - A", "5th Sem - B")),
]
ws_lst = None

# ---------------------------------------------------------------- Mapping & settings
m = ws_map
m.sheet_properties.tabColor = "7030A0"
m.merge_cells("A1:Q1"); hdr(m["A1"], "Settings & Rubric → PO / PSO Mapping Matrix", size=13)
settings = [
 ("Hackathon title", "SDG Hackathon 2026", None),
 ("Institution", "Canara Engineering College, Mangaluru", None),
 ("Department", "Artificial Intelligence and Machine Learning", None),
 ("Academic year", "2026-27", None),
 ("Target score per team (% of maximum marks)", 0.60, "0%"),
 ("Attainment Level 3 if % of teams at/above target ≥", 0.70, "0%"),
 ("Attainment Level 2 if % of teams at/above target ≥", 0.60, "0%"),
 ("Attainment Level 1 if % of teams at/above target ≥", 0.50, "0%"),
 ("Teams shortlisted per round (fixed at 5 by event format)", 5, "0"),
 ("SDG contribution counts as reached at strength ≥ (1–3)", 2, "0"),
]
for i, (lab, val, fmt) in enumerate(settings):
    r = 3 + i
    m.merge_cells(f"A{r}:B{r}")
    calc(m[f"A{r}"], lab, align=LEFT, bold=True)
    m[f"B{r}"].border = BOX
    if i == 8:
        calc(m[f"C{r}"], val, fmt)
    else:
        inp(m[f"C{r}"], val); 
        if fmt: m[f"C{r}"].number_format = fmt
    m.merge_cells(f"C{r}:F{r}")
S_TITLE, S_INST, S_DEPT, S_AY = "Mapping!$C$3", "Mapping!$C$4", "Mapping!$C$5", "Mapping!$C$6"
S_TGT, S_L3, S_L2, S_L1, S_SL = "Mapping!$C$7", "Mapping!$C$8", "Mapping!$C$9", "Mapping!$C$10", "Mapping!$C$11"
S_SDGT = "Mapping!$C$12"
m["H3"] = "Yellow cells are editable. Correlation strength: 3 = High, 2 = Medium, 1 = Low, blank/0 = not mapped."
m["H3"].font = font(italic=True, size=9); m.merge_cells("H3:Q4"); m["H3"].alignment = LEFT
m["H5"] = ("Attainment method: each team's outcome attainment = Σ(criterion score ÷ criterion max × correlation) ÷ Σ correlation. "
           "Event attainment level is set by the % of teams reaching the target score for that outcome. "
           "SDGs are not mapped here: evaluators rate each claimed SDG 0–3 on the event sheets and SDG attainment is the average strength.")
m["H5"].font = font(italic=True, size=9); m.merge_cells("H5:Q8"); m["H5"].alignment = LEFT

MR0 = 15  # first criterion row in matrix
hdr(m["A14"], "#"); hdr(m["B14"], "Rubric Criterion"); hdr(m["C14"], "Max Marks")
for j, o in enumerate(OUTCOMES):
    hdr(m.cell(14, 4 + j), o, bg=("375623" if o == "SDG" else ("7030A0" if o.startswith("PSO") else NAVY)))
for k in range(NC):
    r = MR0 + k
    calc(m.cell(r, 1), k + 1)
    calc(m.cell(r, 2), f"=Rubric!B{4+k}", align=LEFT)
    calc(m.cell(r, 3), f"=Rubric!D{4+k}")
    for j in range(NO):
        c = m.cell(r, 4 + j); inp(c, MAP[k][j] if MAP[k][j] else None)
MRL = MR0 + NC - 1
calc(m.cell(MRL + 1, 2), "Σ Correlation", bold=True, align=LEFT)
calc(m.cell(MRL + 2, 2), "No. of criteria mapped", bold=True, align=LEFT)
for j in range(NO):
    col = L(4 + j)
    calc(m.cell(MRL + 1, 4 + j), f"=SUM({col}{MR0}:{col}{MRL})", bg=AVGF, bold=True)
    calc(m.cell(MRL + 2, 4 + j), f"=COUNTIF({col}{MR0}:{col}{MRL},\">0\")", bg=AVGF)
for c in (1, 3): m.cell(MRL+1, c).border = BOX; m.cell(MRL+2, c).border = BOX
dvm = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=True,
                     error="Correlation must be 0, 1, 2 or 3", errorTitle="Invalid correlation", showErrorMessage=True)
m.add_data_validation(dvm); dvm.add(f"D{MR0}:{L(3+NO)}{MRL}")
# outcome definitions
r0 = MRL + 5
m.merge_cells(f"A{r0-1}:Q{r0-1}"); hdr(m[f"A{r0-1}"], "Programme Outcomes and Programme Specific Outcomes")
for i, (code, desc) in enumerate(POS):
    r = r0 + i
    calc(m.cell(r, 1), code, bold=True); m.merge_cells(start_row=r, start_column=2, end_row=r, end_column=17)
    calc(m.cell(r, 2), desc, align=LEFT)
    m.row_dimensions[r].height = 28 if len(desc) > 90 else 16
m.column_dimensions["A"].width = 7; m.column_dimensions["B"].width = 48; m.column_dimensions["C"].width = 10
for j in range(NO): m.column_dimensions[L(4 + j)].width = 7
m.row_dimensions[14].height = 22
m.freeze_panes = "D15"

# ---------------------------------------------------------------- Rubric
ru = ws_rub
ru.sheet_properties.tabColor = "C00000"
ru.merge_cells("A1:J1"); ru["A1"] = f'={S_TITLE}&" — Evaluation Rubric (for Evaluators)"'
hdr(ru["A1"], ru["A1"].value, size=13)
ru.merge_cells("A2:J2")
ru["A2"] = "Score each criterion out of its Max Marks (decimals allowed). Use the descriptors to place the team in a band, then pick a mark inside that band's range."
ru["A2"].font = font(italic=True, size=9); ru["A2"].alignment = LEFT
heads = ["#", "Criterion", "Short Name", "Max Marks", "Excellent\n(90–100%)", "Good\n(70–89%)",
         "Satisfactory\n(50–69%)", "Needs Improvement\n(<50%)", "Mark bands (auto)", "Mapped PO / PSO (auto, correlation)"]
for i, h in enumerate(heads): hdr(ru.cell(3, 1 + i), h)
for k, (name, short, mx, ex, gd, sa, ni) in enumerate(CRIT):
    r = 4 + k
    calc(ru.cell(r, 1), k + 1)
    inp(ru.cell(r, 2), name, align=LEFT); inp(ru.cell(r, 3), short, align=LEFT); inp(ru.cell(r, 4), mx)
    for i, t in enumerate((ex, gd, sa, ni)): inp(ru.cell(r, 5 + i), t, align=LEFT)
    D = f"D{r}"
    def band(lo, hi): return f'IF({lo}>={hi},{hi}&"",{lo}&"–"&{hi})'
    u9, u7, u5 = f"ROUNDUP(0.9*{D},0)", f"ROUNDUP(0.7*{D},0)", f"ROUNDUP(0.5*{D},0)"
    calc(ru.cell(r, 9),
         f'="E: "&{band(u9, D)}&CHAR(10)&"G: "&{band(u7, f"({u9}-1)")}'
         f'&CHAR(10)&"S: "&{band(u5, f"({u7}-1)")}&CHAR(10)&"NI: "&{band(0, f"({u5}-1)")}',
         align=LEFT, bg=AVGF)
    parts = []
    for j, o in enumerate(OUTCOMES):
        mc = f"Mapping!{L(4+j)}{MR0+k}"
        parts.append(f'IF(N({mc})>0,"{o}("&{mc}&") ","")')
    calc(ru.cell(r, 10), "=" + "&".join(parts), align=LEFT, bg=AVGF)
    ru.row_dimensions[r].height = 64
calc(ru.cell(4 + NC, 3), "Total", bold=True)
calc(ru.cell(4 + NC, 4), f"=SUM(D4:D{3+NC})", bold=True, bg=AVGF)
for w, c in zip([5, 30, 20, 8, 34, 28, 26, 24, 13, 26], "ABCDEFGHIJ"): ru.column_dimensions[c].width = w
ru.row_dimensions[3].height = 32
ru.freeze_panes = "C4"
ru.page_setup.orientation = "landscape"; ru.page_setup.fitToWidth = 1; ru.page_setup.fitToHeight = 0
ru.sheet_properties.pageSetUpPr.fitToPage = True
RUB_TOTAL = f"Rubric!$D${4+NC}"
dvt = DataValidation(type="whole", operator="between", formula1="1", formula2="3", showErrorMessage=True,
                     errorTitle="Invalid strength", error="Enter 1, 2 or 3.")
m.add_data_validation(dvt); dvt.add("C12")

# SDG contribution-strength scale (rated separately from the rubric marks)
STRENGTHS = [(3, "High", "Directly advances a named SDG target (e.g. 6.1), with a measurable indicator and a working solution that plausibly moves it."),
             (2, "Medium", "Meaningfully supports the goal. The goal is clearly served, but the impact is described qualitatively or the target is only broadly named."),
             (1, "Low", "Indirect or enabling. The solution helps something else that in turn helps the goal, such as general awareness."),
             (0, "None", "Claimed but not shown. The SDG label does not match what the solution actually does.")]
sr = 4 + NC + 3
ru.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=10)
hdr(ru.cell(sr, 1), "SDG contribution strength — rate each SDG the team claims (primary and secondary) separately. Not part of the 100 marks.", bg="375623")
hdr(ru.cell(sr + 1, 1), "Rating", bg=SDGF, fg="000000"); hdr(ru.cell(sr + 1, 2), "Strength", bg=SDGF, fg="000000")
ru.merge_cells(start_row=sr + 1, start_column=3, end_row=sr + 1, end_column=10)
hdr(ru.cell(sr + 1, 3), "What the evaluator should see", bg=SDGF, fg="000000")
for i, (v, lab, desc) in enumerate(STRENGTHS):
    r = sr + 2 + i
    calc(ru.cell(r, 1), v, bold=True); calc(ru.cell(r, 2), lab, bold=True, align=LEFT)
    ru.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10); calc(ru.cell(r, 3), desc, align=LEFT)
    for c in range(4, 11): ru.cell(r, c).border = BOX
    ru.row_dimensions[r].height = 30
ru.merge_cells(start_row=sr + 6, start_column=1, end_row=sr + 6, end_column=10)
ru.cell(sr + 6, 1).value = ("Rate the solution's contribution to that goal, not the overall quality of the project. "
                            "Solution strength = average across evaluators; goal attainment = average over all solutions addressing the goal.")
ru.cell(sr + 6, 1).font = font(italic=True, size=9); ru.cell(sr + 6, 1).alignment = LEFT

# ---------------------------------------------------------------- SDG list sheet (at end)
ws_lst = wb.create_sheet("SDG List")
hdr(ws_lst["A1"], "Sustainable Development Goals (dropdown source)")
for i, s in enumerate(SDGS): calc(ws_lst.cell(2 + i, 1), s, align=LEFT)
ws_lst.column_dimensions["A"].width = 52
wb.defined_names["SDGList"] = DefinedName("SDGList", attr_text="'SDG List'!$A$2:$A$18")

# ---------------------------------------------------------------- event sheets
from openpyxl.chart import Series
C_TD = 8                              # team-detail columns A..H
EW = NC + 2                           # evaluator block: 9 criteria + primary & secondary SDG strength
C_E = [9, 9 + EW, 9 + 2 * EW]         # evaluator block starts
C_AVG = 9 + 3 * EW                    # consolidated criterion averages
C_SP, C_SS = C_AVG + NC, C_AVG + NC + 1   # consolidated SDG strengths
C_TOT, C_PCT, C_RANK, C_STAT = C_SS + 1, C_SS + 2, C_SS + 3, C_SS + 4
C_OUT = C_STAT + 1                    # PO1..PSO2
C_KEY = C_OUT + NO                    # hidden rank key
FIRST = 7
info = {}

DEMO_NAMES = ["AquaSense", "AgriBot", "MediTrack", "EduBridge", "SheLead", "PureFlow", "SolarGrid",
              "SkillHub", "SmartFactory", "EqualAccess", "UrbanPulse", "ReCycle AI", "CarbonLens",
              "OceanWatch", "GreenCanopy", "SafeCity", "FoodSaver", "WellMind", "CropDoc", "TrafficIQ"]

def event_sheet(sheet, title, kind, src):
    ws = wb.create_sheet(sheet)
    ws.sheet_properties.tabColor = "2E75B6" if kind == "intra" else "C55A11"
    n = 20 if kind == "intra" else 10
    last = FIRST + n - 1
    lastcol = C_OUT + NO - 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=C_AVG - 1)
    hdr(ws.cell(1, 1), f'={S_TITLE}&" — {title}"', size=13)
    ws.merge_cells(start_row=1, start_column=C_AVG, end_row=1, end_column=lastcol)
    hdr(ws.cell(1, C_AVG), f'={S_INST}&" | Dept. of "&{S_DEPT}&" | AY "&{S_AY}', size=10)
    calc(ws["A2"], "Date:", bold=True); ws.merge_cells("B2:C2"); inp(ws["B2"]); ws["C2"].border = BOX
    calc(ws["D2"], "Venue / Coordinator:", bold=True, align=LEFT); ws.merge_cells("E2:H2"); inp(ws["E2"], align=LEFT)
    for c in "FGH": ws[f"{c}2"].border = BOX
    calc(ws.cell(2, C_AVG), f'="PO/PSO target per team: "&TEXT({S_TGT},"0%")&"  |  SDG reached at strength ≥ "&{S_SDGT}&"  |  Shortlist: top "&{S_SL}', bold=True, align=LEFT)
    ws.merge_cells(start_row=2, start_column=C_AVG, end_row=2, end_column=C_AVG + 12)
    calc(ws["D3"], "Evaluator names →", bold=True, align=LEFT)
    for e, cs in enumerate(C_E):
        calc(ws.cell(3, cs), f"Evaluator {e+1}:", bold=True, bg=E_FILLS[e])
        ws.merge_cells(start_row=3, start_column=cs, end_row=3, end_column=cs + 1)
        ws.merge_cells(start_row=3, start_column=cs + 2, end_row=3, end_column=cs + EW - 1)
        inp(ws.cell(3, cs + 2), align=LEFT)
        for cc in range(cs + 3, cs + EW): ws.cell(3, cc).border = BOX
    groups = [(1, C_TD, "Team Details (Coordinator)" + (" — auto-filled from section shortlists" if kind == "inter" else ""), NAVY, WHITE)]
    for e, cs in enumerate(C_E):
        groups.append((cs, cs + EW - 1, f"Evaluator {e+1} — rubric marks, then SDG contribution strength (0–3)", E_FILLS[e], "000000"))
    groups += [(C_AVG, C_AVG + NC - 1, "Consolidated rubric score (average of evaluators)", "808080", WHITE),
               (C_SP, C_SS, "SDG strength (avg, 0–3)", "375623", WHITE),
               (C_TOT, C_STAT, "Result", "2F5597", WHITE),
               (C_OUT, C_OUT + 10, "Programme Outcome attainment per team (%)", NAVY, WHITE),
               (C_OUT + 11, C_OUT + 12, "PSO attainment (%)", "7030A0", WHITE)]
    for a, b, t, bg, fg in groups:
        if b > a: ws.merge_cells(start_row=4, start_column=a, end_row=4, end_column=b)
        hdr(ws.cell(4, a), t, bg=bg, fg=fg)
        for cc in range(a + 1, b + 1): ws.cell(4, cc).border = BOX
    tds = ["S.No", "Section", "Team ID", "Team Name", "Team Members (USN / Name)", "Primary SDG", "Secondary SDG (optional)", "Problem Statement / Project Title"]
    for i, t in enumerate(tds):
        ws.merge_cells(start_row=5, start_column=i + 1, end_row=6, end_column=i + 1)
        hdr(ws.cell(5, i + 1), t); ws.cell(6, i + 1).border = BOX
    for blk, bg in [(C_E[0], E_FILLS[0]), (C_E[1], E_FILLS[1]), (C_E[2], E_FILLS[2]), (C_AVG, AVGF)]:
        for k in range(NC):
            hdr(ws.cell(5, blk + k), f"=Rubric!C{4+k}", bg=bg, fg="000000", size=8)
            hdr(ws.cell(6, blk + k), f"=Rubric!D{4+k}", bg=bg, fg="000000", bold=True, size=9)
    for cs in C_E + [C_AVG]:
        for off, lab in ((NC, "Primary SDG strength"), (NC + 1, "Secondary SDG strength")):
            hdr(ws.cell(5, cs + off), lab, bg=SDGF, fg="000000", size=8)
            hdr(ws.cell(6, cs + off), "0–3", bg=SDGF, fg="000000", bold=True, size=9)
    for c, t in [(C_TOT, "Total"), (C_PCT, "Score %"), (C_RANK, "Rank"), (C_STAT, "Status")]:
        hdr(ws.cell(5, c), t, bg=RESF, fg="000000")
    hdr(ws.cell(6, C_TOT), f"={RUB_TOTAL}", bg=RESF, fg="000000")
    for c in (C_PCT, C_RANK, C_STAT): hdr(ws.cell(6, c), "", bg=RESF)
    for j, o in enumerate(OUTCOMES):
        hdr(ws.cell(5, C_OUT + j), o, bg=ATTF, fg="000000")
        hdr(ws.cell(6, C_OUT + j), "%", bg=ATTF, fg="000000", size=8)
    ws.row_dimensions[5].height = 48; ws.row_dimensions[4].height = 30
    # ---------------- team rows
    for i in range(n):
        r = FIRST + i
        calc(ws.cell(r, 1), i + 1)
        if kind == "intra":
            inp(ws.cell(r, 2), src)
            for c in range(3, 9): inp(ws.cell(r, c), align=LEFT if c in (4, 5, 8) else CENTER)
        else:
            sA, sB, labA, labB = src
            ss, lab, rk = (sA, labA, i + 1) if i < 5 else (sB, labB, i - 4)
            inp(ws.cell(r, 2), lab)
            for c in range(3, 9):
                colL = L(c)
                f = f"=IFERROR(INDEX('{ss}'!${colL}${FIRST}:${colL}${FIRST+19},MATCH({rk},'{ss}'!${L(C_RANK)}${FIRST}:${L(C_RANK)}${FIRST+19},0))&\"\",\"\")"
                inp(ws.cell(r, c), f, align=LEFT if c in (4, 5, 8) else CENTER)
        for cs in C_E:
            for k in range(NC + 2): inp(ws.cell(r, cs + k))
        for k in range(NC):
            refs = ",".join(f"{L(cs+k)}{r}" for cs in C_E)
            calc(ws.cell(r, C_AVG + k), f'=IF(COUNT({refs})=0,"",AVERAGE({refs}))', "0.0", AVGF)
        pref = ",".join(f"{L(cs+NC)}{r}" for cs in C_E)
        sref = ",".join(f"{L(cs+NC+1)}{r}" for cs in C_E)
        calc(ws.cell(r, C_SP), f'=IF(OR($F{r}="",COUNT({pref})=0),"",AVERAGE({pref}))', "0.00", SDGF, True)
        calc(ws.cell(r, C_SS), f'=IF(OR($G{r}="",$G{r}=$F{r},COUNT({sref})=0),"",AVERAGE({sref}))', "0.00", SDGF, True)
        a0, a1 = L(C_AVG), L(C_AVG + NC - 1)
        T = f"{L(C_TOT)}{r}"
        calc(ws.cell(r, C_TOT), f'=IF(COUNT({a0}{r}:{a1}{r})=0,"",SUM({a0}{r}:{a1}{r}))', "0.0", RESF, True)
        calc(ws.cell(r, C_PCT), f'=IF({T}="","",{T}/{L(C_TOT)}$6)', "0.0%", RESF)
        K = L(C_KEY)
        calc(ws.cell(r, C_KEY), f'=IF({T}="","",{T}+N({L(C_AVG+3)}{r})/10000)', "0.0000")
        calc(ws.cell(r, C_RANK), f'=IF({K}{r}="","",COUNTIF(${K}${FIRST}:${K}${last},">"&{K}{r})+COUNTIF(${K}${FIRST}:{K}{r},{K}{r}))', "0", RESF, True)
        calc(ws.cell(r, C_STAT), f'=IF({L(C_RANK)}{r}="","",IF({L(C_RANK)}{r}<={S_SL},"Shortlisted","Not shortlisted"))', None, RESF)
        for j in range(NO):
            mc = L(4 + j)
            terms = "+".join(f"{L(C_AVG+k)}{r}/{L(C_AVG+k)}$6*Mapping!${mc}${MR0+k}" for k in range(NC))
            calc(ws.cell(r, C_OUT + j), f"=IFERROR(({terms})/SUM(Mapping!${mc}${MR0}:${mc}${MRL}),\"\")", "0.0%", ATTF)
    ws.column_dimensions[L(C_KEY)].hidden = True
    # ---------------- validation
    for cs in C_E:
        for k in range(NC):
            colL = L(cs + k)
            dv = DataValidation(type="decimal", operator="between", formula1="0", formula2=f"${colL}$6",
                                showErrorMessage=True, errorTitle="Score out of range",
                                error="Enter a score between 0 and the criterion's Max Marks (row 6).")
            ws.add_data_validation(dv); dv.add(f"{colL}{FIRST}:{colL}{last}")
        dvs = DataValidation(type="whole", operator="between", formula1="0", formula2="3", showErrorMessage=True,
                             errorTitle="SDG strength", error="Enter 0 (None), 1 (Low), 2 (Medium) or 3 (High).",
                             showInputMessage=True, promptTitle="SDG contribution strength",
                             prompt="3 High: advances a named target with a measurable indicator. 2 Medium: clearly supports the goal. 1 Low: indirect. 0 None: claimed only.")
        ws.add_data_validation(dvs); dvs.add(f"{L(cs+NC)}{FIRST}:{L(cs+NC+1)}{last}")
    dvl = DataValidation(type="list", formula1="SDGList", allow_blank=True, showErrorMessage=True,
                         errorTitle="Pick an SDG", error="Select an SDG from the list.")
    ws.add_data_validation(dvl); dvl.add(f"F{FIRST}:G{last}")
    stat = f"{L(C_STAT)}{FIRST}:{L(C_STAT)}{last}"
    ws.conditional_formatting.add(stat, CellIsRule(operator="equal", formula=['"Shortlisted"'], fill=fill("C6EFCE"), font=Font(name=F, color="006100", bold=True)))
    # grey out secondary-strength inputs when the team has no (distinct) secondary SDG
    for cs in C_E:
        rng = f"{L(cs+NC+1)}{FIRST}:{L(cs+NC+1)}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR($G{FIRST}="",$G{FIRST}=$F{FIRST})'], fill=fill("D9D9D9")))
    # ---------------- summary
    s = last + 2
    rngT = lambda col: f"{L(col)}${FIRST}:{L(col)}${last}"
    labels = ["Average score (marks)", "Average score (% of max)", "Average attainment (%)",
              "% of teams at/above target", "Attainment level (0–3)", "Teams evaluated",
              "SDG attainment: average contribution strength (0–3)",
              "SDG: % of rated team–SDG pairs at/above target strength",
              "SDG: team–SDG pairs rated", "SDG: team–SDG pairs claimed"]
    for i, t in enumerate(labels):
        calc(ws.cell(s + i, 4), t, bold=True, align=LEFT, bg=SDGF if t.startswith("SDG") else RESF)
    for k in range(NC):
        c = C_AVG + k
        calc(ws.cell(s, c), f'=IFERROR(AVERAGE({rngT(c)}),"")', "0.0", AVGF)
        calc(ws.cell(s + 1, c), f'=IFERROR({L(c)}{s}/{L(c)}$6,"")', "0.0%", AVGF)
    for c in (C_SP, C_SS):
        calc(ws.cell(s, c), f'=IFERROR(AVERAGE({rngT(c)}),"")', "0.00", SDGF, True)
    calc(ws.cell(s, C_TOT), f'=IFERROR(AVERAGE({rngT(C_TOT)}),"")', "0.0", RESF, True)
    calc(ws.cell(s + 1, C_TOT), f'=IFERROR({L(C_TOT)}{s}/{L(C_TOT)}$6,"")', "0.0%", RESF, True)
    for j in range(NO):
        c = C_OUT + j; col = L(c)
        calc(ws.cell(s + 2, c), f'=IFERROR(AVERAGE({rngT(c)}),"")', "0.0%", ATTF, True)
        calc(ws.cell(s + 3, c), f'=IFERROR(COUNTIF({rngT(c)},">="&{S_TGT})/COUNT({rngT(c)}),"")', "0.0%", ATTF, True)
        p = f"{col}{s+3}"
        calc(ws.cell(s + 4, c), f'=IF({p}="","",IF({p}>={S_L3},3,IF({p}>={S_L2},2,IF({p}>={S_L1},1,0))))', "0", ATTF, True)
    SPr, SSr = rngT(C_SP), rngT(C_SS)
    Fr, Gr = f"$F${FIRST}:$F${last}", f"$G${FIRST}:$G${last}"
    calc(ws.cell(s + 5, 5), f"=COUNT({rngT(C_TOT)})", "0", bold=True)
    calc(ws.cell(s + 6, 5), f'=IFERROR((SUM({SPr})+SUM({SSr}))/(COUNT({SPr})+COUNT({SSr})),"")', "0.00", SDGF, True)
    calc(ws.cell(s + 7, 5), f'=IFERROR((COUNTIF({SPr},">="&{S_SDGT})+COUNTIF({SSr},">="&{S_SDGT}))/(COUNT({SPr})+COUNT({SSr})),"")', "0.0%", SDGF, True)
    calc(ws.cell(s + 8, 5), f"=COUNT({SPr})+COUNT({SSr})", "0", SDGF, True)
    calc(ws.cell(s + 9, 5), f'=COUNTIF({Fr},"?*")+SUMPRODUCT(--({Gr}<>""),--({Gr}<>{Fr}))', "0", SDGF, True)
    lvl = f"{L(C_OUT)}{s+4}:{L(C_OUT+NO-1)}{s+4}"
    for v, colr in [(3, "C6EFCE"), (2, "FFEB9C"), (1, "FCE4D6"), (0, "FFC7CE")]:
        ws.conditional_formatting.add(lvl, FormulaRule(formula=[f'AND(ISNUMBER({L(C_OUT)}{s+4}),{L(C_OUT)}{s+4}={v})'], fill=fill(colr)))
    # ---------------- SDG goal-wise table
    g0 = s + 12
    ws.merge_cells(start_row=g0 - 1, start_column=4, end_row=g0 - 1, end_column=8)
    hdr(ws.cell(g0 - 1, 4), "SDG goal-wise attainment (from contribution-strength ratings)", bg="375623")
    for i, t in enumerate(["Sustainable Development Goal", "Solutions (teams claiming the goal)", "Solutions rated",
                           "Avg contribution strength (0–3)", "% rated at/above target strength"]):
        hdr(ws.cell(g0, 4 + i), t, bg=SDGF, fg="000000")
    ws.row_dimensions[g0].height = 45
    for i, g in enumerate(SDGS):
        r = g0 + 1 + i
        calc(ws.cell(r, 4), f"='SDG List'!A{2+i}", align=LEFT)
        D = f"$D{r}"
        calc(ws.cell(r, 5), f"=COUNTIF({Fr},{D})+COUNTIF({Gr},{D})-COUNTIFS({Fr},{D},{Gr},{D})", "0")
        calc(ws.cell(r, 6), f'=COUNTIFS({Fr},{D},{SPr},">=0")+COUNTIFS({Gr},{D},{SSr},">=0")', "0")
        calc(ws.cell(r, 7), f'=IFERROR((SUMIF({Fr},{D},{SPr})+SUMIF({Gr},{D},{SSr}))/F{r},"")', "0.00", bold=True)
        calc(ws.cell(r, 8), f'=IFERROR((COUNTIFS({Fr},{D},{SPr},">="&{S_SDGT})+COUNTIFS({Gr},{D},{SSr},">="&{S_SDGT}))/F{r},"")', "0.0%")
    gl = g0 + len(SDGS)
    # ---------------- layout
    for c, w in zip("ABCDEFGH", [5, 12, 9, 30, 30, 26, 26, 34]): ws.column_dimensions[c].width = w
    for c in range(9, C_KEY): ws.column_dimensions[L(c)].width = 8.5
    ws.column_dimensions[L(C_STAT)].width = 14
    for r in range(FIRST, last + 1): ws.row_dimensions[r].height = 30
    ws.freeze_panes = "E7"
    ws.protection.sheet = True
    ws.protection.formatColumns = False; ws.protection.formatRows = False
    ws.protection.selectLockedCells = False; ws.protection.selectUnlockedCells = False
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    # ---------------- charts
    cats_out = Reference(ws, min_col=C_OUT, max_col=C_OUT + NO - 1, min_row=5)
    ch = BarChart(); ch.type = "col"; ch.title = "PO / PSO attainment"
    for rr, nm in [(s + 2, "Average attainment"), (s + 3, "% teams ≥ target")]:
        ch.series.append(Series(Reference(ws, min_col=C_OUT, max_col=C_OUT + NO - 1, min_row=rr), title=nm))
    ch.set_categories(cats_out); ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
    ch.y_axis.number_format = "0%"; ch.legend.position = "b"; ch.width, ch.height = 20, 9; unhide_axes(ch)
    ws.add_chart(ch, f"I{g0-1}")
    ch2 = BarChart(); ch2.type = "col"; ch2.title = "PO / PSO attainment level (0–3)"
    ch2.series.append(Series(Reference(ws, min_col=C_OUT, max_col=C_OUT + NO - 1, min_row=s + 4), title="Level"))
    ch2.set_categories(cats_out); ch2.y_axis.scaling.min = 0; ch2.y_axis.scaling.max = 3; ch2.y_axis.majorUnit = 1
    ch2.legend = None; ch2.width, ch2.height = 20, 9; unhide_axes(ch2); datalabels(ch2)
    ws.add_chart(ch2, f"W{g0-1}")
    ch5 = BarChart(); ch5.type = "bar"; ch5.title = "SDG goal-wise average contribution strength (0–3)"
    ch5.series.append(Series(Reference(ws, min_col=7, min_row=g0 + 1, max_row=gl), title="Avg strength"))
    ch5.set_categories(Reference(ws, min_col=4, min_row=g0 + 1, max_row=gl))
    ch5.x_axis.scaling.orientation = "maxMin"; ch5.y_axis.scaling.min = 0; ch5.y_axis.scaling.max = 3; ch5.y_axis.majorUnit = 1
    ch5.legend = None; ch5.width, ch5.height = 20, 11; unhide_axes(ch5)
    ws.add_chart(ch5, f"I{g0+18}")
    ch4 = BarChart(); ch4.type = "col"; ch4.title = "Criterion-wise average score (% of max)"
    ch4.series.append(Series(Reference(ws, min_col=C_AVG, max_col=C_AVG + NC - 1, min_row=s + 1), title="Avg %"))
    ch4.set_categories(Reference(ws, min_col=C_AVG, max_col=C_AVG + NC - 1, min_row=5))
    ch4.y_axis.scaling.min = 0; ch4.y_axis.scaling.max = 1; ch4.y_axis.number_format = "0%"
    ch4.legend = None; ch4.width, ch4.height = 20, 11; unhide_axes(ch4)
    ws.add_chart(ch4, f"W{g0+18}")
    ch3 = BarChart(); ch3.type = "bar"; ch3.title = "Team total scores"
    ch3.series.append(Series(Reference(ws, min_col=C_TOT, min_row=FIRST, max_row=last), title="Total"))
    ch3.set_categories(Reference(ws, min_col=4, min_row=FIRST, max_row=last))
    ch3.x_axis.scaling.orientation = "maxMin"; ch3.y_axis.scaling.min = 0; ch3.y_axis.scaling.max = 100
    ch3.legend = None; ch3.width, ch3.height = 20, 11 if kind == "intra" else 8; unhide_axes(ch3); datalabels(ch3)
    ws.add_chart(ch3, f"I{g0+41}")
    info[sheet] = dict(s=s, g0=g0, last=last, n=n, title=title, kind=kind)
    return ws

def datalabels(chart):
    chart.dataLabels = DataLabelList(); chart.dataLabels.showVal = True
    chart.dataLabels.showSerName = False; chart.dataLabels.showCatName = False
    chart.dataLabels.showLegendKey = False; chart.dataLabels.showPercent = False

for ev in EVENTS: event_sheet(*ev)
wb.move_sheet("SDG List", offset=len(wb.sheetnames))

# ---------------------------------------------------------------- Dashboard
d = ws_dash
d.sheet_properties.tabColor = "00B050"
LASTC = 2 + NO
d.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LASTC); hdr(d["A1"], f'={S_TITLE}&" — Attainment Dashboard (PO / PSO / SDG)"', size=14)
d.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LASTC)
d["A2"] = f'={S_INST}&" | Dept. of "&{S_DEPT}&" | AY "&{S_AY}&" | PO/PSO target per team: "&TEXT({S_TGT},"0%")&" | SDG reached at strength ≥ "&{S_SDGT}'
d["A2"].font = font(italic=True, size=10); d["A2"].alignment = CENTER
EV_LABEL = {"3A Intra": "3rd Sem A — Intra", "3B Intra": "3rd Sem B — Intra", "3 Inter": "3rd Sem — Inter-Section",
            "5A Intra": "5th Sem A — Intra", "5B Intra": "5th Sem B — Intra", "5 Inter": "5th Sem — Inter-Section"}

def dash_table(r0, title, srow_off, fmt):
    d.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=LASTC)
    hdr(d.cell(r0, 1), title)
    hdr(d.cell(r0 + 1, 1), "Event", bg=RESF, fg="000000"); hdr(d.cell(r0 + 1, 2), "Teams evaluated", bg=RESF, fg="000000")
    for j, o in enumerate(OUTCOMES): hdr(d.cell(r0 + 1, 3 + j), o, bg=ATTF, fg="000000")
    for i, (sh, *_rest) in enumerate(EVENTS):
        r = r0 + 2 + i; I = info[sh]
        calc(d.cell(r, 1), EV_LABEL[sh], align=LEFT)
        calc(d.cell(r, 2), f"='{sh}'!E{I['s']+5}", "0")
        for j in range(NO):
            calc(d.cell(r, 3 + j), f"='{sh}'!{L(C_OUT+j)}{I['s']+srow_off}", fmt)
    return r0 + 2, r0 + 7
t1a, t1b = dash_table(4, "Table 1 — PO / PSO average attainment (%) by event", 2, "0.0%")
t2a, t2b = dash_table(14, "Table 2 — PO / PSO: % of teams at/above target by event", 3, "0.0%")
t3a, t3b = dash_table(24, "Table 3 — PO / PSO attainment level (0–3) by event", 4, "0")
for (a, b) in [(t1a, t1b), (t2a, t2b)]:
    r = b + 1
    calc(d.cell(r, 1), "Overall (weighted by teams)", bold=True, bg=RESF, align=LEFT)
    calc(d.cell(r, 2), f"=SUM(B{a}:B{b})", "0", bold=True, bg=RESF)
    for j in range(NO):
        c = L(3 + j)
        calc(d.cell(r, 3 + j), f'=IFERROR(SUMPRODUCT({c}{a}:{c}{b},$B${a}:$B${b})/SUMPRODUCT(--ISNUMBER({c}{a}:{c}{b}),$B${a}:$B${b}),"")',
             "0.0%", bg=RESF, bold=True)
r = t3b + 1
calc(d.cell(r, 1), "Overall attainment level", bold=True, bg=RESF, align=LEFT)
calc(d.cell(r, 2), f"=B{t2b+1}", "0", bold=True, bg=RESF)
for j in range(NO):
    p = f"{L(3+j)}{t2b+1}"
    calc(d.cell(r, 3 + j), f'=IF({p}="","",IF({p}>={S_L3},3,IF({p}>={S_L2},2,IF({p}>={S_L1},1,0))))', "0", bg=RESF, bold=True)
lv = f"C{t3a}:{L(LASTC)}{t3b+1}"
for v, colr in [(3, "C6EFCE"), (2, "FFEB9C"), (1, "FCE4D6"), (0, "FFC7CE")]:
    d.conditional_formatting.add(lv, FormulaRule(formula=[f"AND(ISNUMBER(C{t3a}),C{t3a}={v})"], fill=fill(colr)))
OVR_AVG, OVR_PCT, OVR_LVL = t1b + 1, t2b + 1, t3b + 1

# Table 4 — SDG attainment by event
p0 = 34
d.merge_cells(start_row=p0, start_column=1, end_row=p0, end_column=5)
hdr(d.cell(p0, 1), "Table 4 — SDG attainment by event (contribution strength, 0–3)", bg="375623")
for i, t in enumerate(["Event", "Team–SDG pairs claimed", "Pairs rated", "Avg contribution strength (0–3)", "% rated at/above target strength"]):
    hdr(d.cell(p0 + 1, 1 + i), t, bg=SDGF, fg="000000")
d.row_dimensions[p0 + 1].height = 45
for i, (sh, *_x) in enumerate(EVENTS):
    r = p0 + 2 + i; s0 = info[sh]["s"]
    calc(d.cell(r, 1), EV_LABEL[sh], align=LEFT)
    calc(d.cell(r, 2), f"='{sh}'!E{s0+9}", "0")
    calc(d.cell(r, 3), f"='{sh}'!E{s0+8}", "0")
    calc(d.cell(r, 4), f"='{sh}'!E{s0+6}", "0.00", bold=True)
    calc(d.cell(r, 5), f"='{sh}'!E{s0+7}", "0.0%")
pa, pb = p0 + 2, p0 + 7
r = pb + 1
calc(d.cell(r, 1), "Overall (all rated pairs)", bold=True, bg=RESF, align=LEFT)
calc(d.cell(r, 2), f"=SUM(B{pa}:B{pb})", "0", bold=True, bg=RESF)
calc(d.cell(r, 3), f"=SUM(C{pa}:C{pb})", "0", bold=True, bg=RESF)
calc(d.cell(r, 4), f'=IFERROR(SUMPRODUCT(D{pa}:D{pb},C{pa}:C{pb})/C{r},"")', "0.00", bold=True, bg=RESF)
calc(d.cell(r, 5), f'=IFERROR(SUMPRODUCT(E{pa}:E{pb},C{pa}:C{pb})/C{r},"")', "0.0%", bold=True, bg=RESF)
OVR_SDG = r

# Table 5 — SDG goal-wise, all events
q0 = 45
d.merge_cells(start_row=q0, start_column=1, end_row=q0, end_column=11)
hdr(d.cell(q0, 1), "Table 5 — SDG goal-wise attainment, all events (contribution strength, 0–3)", bg="375623")
hs = ["Sustainable Development Goal"] + [f"{EV_LABEL[e[0]]}: solutions" for e in EVENTS] + ["Total solutions", "Solutions rated", "Avg contribution strength (0–3)", "% rated at/above target strength"]
for i, t in enumerate(hs): hdr(d.cell(q0 + 1, 1 + i), t, bg=SDGF, fg="000000")
d.row_dimensions[q0 + 1].height = 58
for i, g in enumerate(SDGS):
    r = q0 + 2 + i
    calc(d.cell(r, 1), f"='SDG List'!A{2+i}", align=LEFT)
    for e, (sh, *_x) in enumerate(EVENTS):
        calc(d.cell(r, 2 + e), f"='{sh}'!E{info[sh]['g0']+1+i}", "0")
    calc(d.cell(r, 8), f"=SUM(B{r}:G{r})", "0", bold=True)
    rated = "+".join(f"'{sh}'!F{info[sh]['g0']+1+i}" for sh, *_x in EVENTS)
    wsum = "+".join(f"N('{sh}'!G{info[sh]['g0']+1+i})*'{sh}'!F{info[sh]['g0']+1+i}" for sh, *_x in EVENTS)
    psum = "+".join(f"N('{sh}'!H{info[sh]['g0']+1+i})*'{sh}'!F{info[sh]['g0']+1+i}" for sh, *_x in EVENTS)
    calc(d.cell(r, 9), f"={rated}", "0")
    calc(d.cell(r, 10), f'=IFERROR(({wsum})/I{r},"")', "0.00", bold=True)
    calc(d.cell(r, 11), f'=IFERROR(({psum})/I{r},"")', "0.0%")
qL = q0 + 1 + len(SDGS)
calc(d.cell(qL + 1, 1), "Total", bold=True, bg=RESF, align=LEFT)
for c in range(2, 10): calc(d.cell(qL + 1, c), f"=SUM({L(c)}{q0+2}:{L(c)}{qL})", "0", bold=True, bg=RESF)
st = f"J{q0+2}:J{qL}"
d.conditional_formatting.add(st, FormulaRule(formula=[f"AND(ISNUMBER(J{q0+2}),J{q0+2}>={S_SDGT})"], fill=fill("C6EFCE")))
d.conditional_formatting.add(st, FormulaRule(formula=[f"AND(ISNUMBER(J{q0+2}),J{q0+2}<{S_SDGT})"], fill=fill("FCE4D6")))

# Table 6 — shortlisted teams
k0 = qL + 4
d.merge_cells(start_row=k0, start_column=1, end_row=k0, end_column=13)
hdr(d.cell(k0, 1), "Table 6 — Shortlisted teams (auto-ranked; ties broken by Technical Implementation score)")
hdr(d.cell(k0 + 1, 1), "Rank", bg=RESF, fg="000000")
for e, (sh, *_x) in enumerate(EVENTS):
    c = 2 + 2 * e
    d.merge_cells(start_row=k0 + 1, start_column=c, end_row=k0 + 1, end_column=c + 1)
    hdr(d.cell(k0 + 1, c), EV_LABEL[sh], bg=RESF, fg="000000"); d.cell(k0 + 1, c + 1).border = BOX
    hdr(d.cell(k0 + 2, c), "Team", bg=AVGF, fg="000000"); hdr(d.cell(k0 + 2, c + 1), "Score", bg=AVGF, fg="000000")
hdr(d.cell(k0 + 2, 1), "", bg=AVGF)
for rk in range(1, 6):
    r = k0 + 2 + rk
    calc(d.cell(r, 1), rk, bold=True)
    for e, (sh, *_x) in enumerate(EVENTS):
        I = info[sh]; c = 2 + 2 * e
        RK = f"'{sh}'!${L(C_RANK)}${FIRST}:${L(C_RANK)}${I['last']}"
        calc(d.cell(r, c), f"=IFERROR(INDEX('{sh}'!$D${FIRST}:$D${I['last']},MATCH($A{r},{RK},0))&\"\",\"\")", align=LEFT)
        calc(d.cell(r, c + 1), f"=IFERROR(INDEX('{sh}'!${L(C_TOT)}${FIRST}:${L(C_TOT)}${I['last']},MATCH($A{r},{RK},0)),\"\")", "0.0")

d.column_dimensions["A"].width = 30
for c in range(2, LASTC + 1): d.column_dimensions[L(c)].width = 11
d.freeze_panes = "B4"
d.protection.sheet = True; d.protection.formatColumns = False; d.protection.formatRows = False

cats = Reference(d, min_col=3, max_col=LASTC, min_row=5)
c1 = BarChart(); c1.type = "col"; c1.title = "Overall PO / PSO attainment"
c1.series.append(Series(Reference(d, min_col=3, max_col=LASTC, min_row=OVR_AVG), title="Average attainment"))
c1.series.append(Series(Reference(d, min_col=3, max_col=LASTC, min_row=OVR_PCT), title="% teams ≥ target"))
c1.set_categories(cats); c1.y_axis.scaling.min = 0; c1.y_axis.scaling.max = 1; c1.y_axis.number_format = "0%"
c1.legend.position = "b"; c1.width, c1.height = 22, 10; unhide_axes(c1)
d.add_chart(c1, "R3")
c2 = BarChart(); c2.type = "col"; c2.title = "Overall PO / PSO attainment level (0–3)"
c2.series.append(Series(Reference(d, min_col=3, max_col=LASTC, min_row=OVR_LVL), title="Level"))
c2.set_categories(cats); c2.y_axis.scaling.min = 0; c2.y_axis.scaling.max = 3; c2.y_axis.majorUnit = 1
c2.legend = None; c2.width, c2.height = 22, 10; unhide_axes(c2); datalabels(c2)
d.add_chart(c2, "R23")
c3 = BarChart(); c3.type = "col"; c3.title = "PO / PSO average attainment by event"
for i in range(6):
    c3.series.append(Series(Reference(d, min_col=3, max_col=LASTC, min_row=t1a + i), title=EV_LABEL[EVENTS[i][0]]))
c3.set_categories(cats); c3.y_axis.scaling.min = 0; c3.y_axis.scaling.max = 1; c3.y_axis.number_format = "0%"
c3.legend.position = "b"; c3.width, c3.height = 22, 10; unhide_axes(c3)
d.add_chart(c3, "AF23")
c4 = RadarChart(); c4.type = "marker"; c4.title = "PO / PSO profile by event (avg %)"
for i in range(6):
    c4.series.append(Series(Reference(d, min_col=3, max_col=LASTC, min_row=t1a + i), title=EV_LABEL[EVENTS[i][0]]))
c4.set_categories(cats); c4.y_axis.scaling.min = 0; c4.y_axis.scaling.max = 1
c4.y_axis.number_format = "0%"; c4.width, c4.height = 22, 10; c4.legend.position = "b"
c4.x_axis.delete = False; c4.y_axis.delete = False
d.add_chart(c4, "AF3")
c7 = BarChart(); c7.type = "col"; c7.title = "SDG attainment by event (avg contribution strength, 0–3)"
c7.series.append(Series(Reference(d, min_col=4, min_row=pa, max_row=pb + 1), title="Avg strength"))
c7.set_categories(Reference(d, min_col=1, min_row=pa, max_row=pb + 1))
c7.y_axis.scaling.min = 0; c7.y_axis.scaling.max = 3; c7.y_axis.majorUnit = 1
c7.legend = None; c7.width, c7.height = 22, 10; unhide_axes(c7); datalabels(c7)
d.add_chart(c7, "R43")
c6 = BarChart(); c6.type = "bar"; c6.title = "SDG goal-wise average contribution strength (0–3)"
c6.series.append(Series(Reference(d, min_col=10, min_row=q0 + 2, max_row=qL), title="Avg strength"))
c6.set_categories(Reference(d, min_col=1, min_row=q0 + 2, max_row=qL))
c6.x_axis.scaling.orientation = "maxMin"; c6.y_axis.scaling.min = 0; c6.y_axis.scaling.max = 3; c6.y_axis.majorUnit = 1
c6.legend = None; c6.width, c6.height = 22, 13; unhide_axes(c6)
d.add_chart(c6, "AF43")
c5 = BarChart(); c5.type = "bar"; c5.title = "SDG coverage (solutions per goal)"; c5.grouping = "stacked"; c5.overlap = 100
for e in range(6):
    c5.series.append(Series(Reference(d, min_col=2 + e, min_row=q0 + 2, max_row=qL), title=EV_LABEL[EVENTS[e][0]]))
c5.set_categories(Reference(d, min_col=1, min_row=q0 + 2, max_row=qL))
c5.x_axis.scaling.orientation = "maxMin"; c5.legend.position = "b"; c5.width, c5.height = 22, 13; unhide_axes(c5)
d.add_chart(c5, "R64")

# ---------------------------------------------------------------- Instructions
ins = ws_ins
ins.sheet_properties.tabColor = "FFC000"
ins.merge_cells("A1:J1"); hdr(ins["A1"], f'={S_TITLE}&" — Evaluation Workbook: How to use"', size=14)
lines = [
 ("Hackathon structure", None),
 ("Round 1", "3rd Sem Section A and Section B intra-class hackathons → top 5 teams from each section shortlisted (sheets '3A Intra', '3B Intra')."),
 ("Round 1", "5th Sem Section A and Section B intra-class hackathons → top 5 teams from each section shortlisted (sheets '5A Intra', '5B Intra')."),
 ("Round 2", "3rd Sem inter-section hackathon: the 10 shortlisted teams are pulled in automatically → top 5 shortlisted (sheet '3 Inter')."),
 ("Round 2", "5th Sem inter-section hackathon: the 10 shortlisted teams are pulled in automatically → top 5 shortlisted (sheet '5 Inter')."),
 ("Who enters what", None),
 ("Coordinator (once)", "'Mapping' sheet: title, department, PO/PSO target %, attainment thresholds, SDG target strength, and (optionally) the rubric → PO/PSO correlation matrix. 'Rubric' sheet: criteria, max marks, descriptors."),
 ("Coordinator (per event)", "In each event sheet: date, venue, evaluator names, and team details (Team ID, name, members, Primary SDG and optional Secondary SDG from the dropdown, problem statement). Intra sheets hold up to 20 teams."),
 ("Evaluators", "In your own coloured block: the 9 rubric marks (0 to each criterion's maximum in row 6), then the contribution strength (0–3) for the team's Primary SDG and, if the team has one, its Secondary SDG. The secondary cell turns grey when there is no secondary SDG. Leave a block blank if fewer than 3 evaluators are judging."),
 ("Everything else is automatic", "Averages across evaluators, total, %, rank, shortlist status, per-team PO/PSO attainment, event attainment levels, SDG strength and goal-wise attainment, the Dashboard and all charts."),
 ("Colour legend", None),
 ("Yellow cell, blue text", "Input cell — type here."),
 ("Grey / lavender / green cells", "Calculated — do not edit (event sheets and Dashboard are protected without a password; Review ▸ Unprotect Sheet if you need to change the layout)."),
 ("PO / PSO attainment", None),
 ("Criterion %", "Average of the evaluators' marks for a criterion ÷ its maximum marks."),
 ("Team outcome attainment", "For each PO/PSO: Σ(criterion % × correlation) ÷ Σ correlation, using the 3/2/1 matrix on the 'Mapping' sheet."),
 ("Event attainment level", "% of teams whose outcome attainment ≥ target (default 60%). Level 3 if ≥70% of teams, Level 2 if ≥60%, Level 1 if ≥50%, else 0 (all editable on 'Mapping')."),
 ("SDG attainment", None),
 ("Contribution strength", "Each evaluator rates how strongly the solution advances each SDG the team claims: 3 High (advances a named target with a measurable indicator), 2 Medium (clearly supports the goal), 1 Low (indirect), 0 None (claimed only). Scale and descriptors are on the 'Rubric' sheet. Not part of the 100 marks."),
 ("Solution strength", "Average of the evaluators' ratings for that team and SDG (0–3)."),
 ("Goal-wise attainment", "For each of the 17 SDGs: average strength of all solutions addressing it (primary or secondary), plus the % of rated solutions at or above the target strength (default 2, Medium)."),
 ("Event SDG attainment", "Average strength over every rated team–SDG pair in the event; the Dashboard pools all events weighted by the number of rated pairs."),
 ("Ranking & ties", "Rank is by total rubric score only; SDG ratings do not change it. Ties are broken by the Technical Implementation (C4) score, then by row order. Inter-section sheets pull rank 1–5 of each section; you can overtype a team cell there if a team is replaced."),
 ("Example row (format only)", None),
]
r = 3
for a, b in lines:
    if b is None:
        ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        hdr(ins.cell(r, 1), a, bg="2F5597"); r += 1; continue
    calc(ins.cell(r, 1), a, bold=True, align=LEFT, bg=AVGF)
    ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ins.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
    calc(ins.cell(r, 3), b, align=LEFT)
    ins.row_dimensions[r].height = 46 if len(b) > 220 else 32 if len(b) > 110 else 18
    if a == "Yellow cell, blue text": inp(ins.cell(r, 3), b, align=LEFT)
    r += 1
ex_h = ["Section", "Team ID", "Team Name", "Primary SDG", "Evaluator 1: C1…C9", "Evaluator 1: SDG strength (primary, secondary)"]
ex_v = ["3rd Sem - A", "3A-07", "AquaSense", "SDG 6 - Clean Water and Sanitation", "8, 7, 12, 16, 8, 7, 7, 4, 8", "3, (blank)"]
for i, (h, v) in enumerate(zip(ex_h, ex_v)):
    hdr(ins.cell(r, 1 + i), h, bg=RESF, fg="000000"); calc(ins.cell(r + 1, 1 + i), v, align=LEFT)
ins.row_dimensions[r].height = 32; ins.row_dimensions[r + 1].height = 30
for c, w in zip("ABCDEFGHIJ", [16, 14, 14, 30, 24, 26, 12, 12, 12, 12]): ins.column_dimensions[c].width = w

# ---------------------------------------------------------------- demo data
if DEMO:
    SDGS_DEMO = [SDGS[i] for i in (2, 3, 5, 6, 8, 10, 11, 12, 1, 4)]
    for sh, t, kind, src in EVENTS:
        ws = wb[sh]; I = info[sh]
        nteams = 16 if kind == "intra" else 10
        ws["B2"] = "15-10-2026"; ws["E2"] = "Seminar Hall / Dr. Coordinator"
        for e, cs in enumerate(C_E): ws.cell(3, cs + 2).value = f"Evaluator {chr(65+e)}"
        names = random.sample(DEMO_NAMES, 16)
        for i in range(nteams):
            r = FIRST + i
            skill = random.uniform(0.45, 0.92)
            has_sec = False
            if kind == "intra":
                ws.cell(r, 3).value = f"{sh[:2]}-{i+1:02d}"; ws.cell(r, 4).value = names[i]
                ws.cell(r, 5).value = "4 members"; ws.cell(r, 6).value = random.choice(SDGS_DEMO)
                if random.random() < 0.3:
                    ws.cell(r, 7).value = random.choice([x for x in SDGS_DEMO if x != ws.cell(r, 6).value]); has_sec = True
                ws.cell(r, 8).value = f"{names[i]} — SDG-focused prototype"
            else:
                skill = random.uniform(0.7, 0.95); has_sec = True
            base = random.choice([1, 2, 2, 3, 3])
            for e, cs in enumerate(C_E):
                if kind == "intra" and e == 2: continue
                for k in range(NC):
                    mx = CRIT[k][2]
                    ws.cell(r, cs + k).value = round(max(0, min(mx, mx * random.gauss(skill, 0.08))) * 2) / 2
                ws.cell(r, cs + NC).value = max(0, min(3, base + random.choice([-1, 0, 0, 1])))
                if has_sec: ws.cell(r, cs + NC + 1).value = max(0, min(3, base - 1 + random.choice([-1, 0, 0, 1])))
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
